from docx import Document
from datetime import datetime
import os

def create_student_word_file(student, save_dir="."):
    """
    Creates a base Word file with student details in the required format.
    No update info is appended. Used for clean upload initialization.
    """
    try:
        file_name = f"{student['Name'].replace(' ', '')}{student['Enrollment']}.docx"
        file_path = os.path.join(save_dir, file_name)

        doc = Document()
        ordered_fields = ["Roll", "Name", "Enrollment", "Password", "Sex", "DSA", "AC", "DC", "MATLAB"]
        full_line = ", ".join([f"{key}: {student[key]}" for key in ordered_fields if key in student])
        doc.add_paragraph(full_line)

        doc.save(file_path)
        print(f"Created base Word file: {file_path}")
    except Exception as e:
        print(f"Error creating base Word file for {student['Enrollment']}: {e}")

def update_or_create_student_word_file(student, subject, new_marks, faculty_name, save_dir="."):
    """
    Creates or updates a Word file for a student in the required format.
    If the file exists, appends the updated subject with timestamp.
    If not, creates a new file with full details and appends update info.

    Args:
        student (dict): Dictionary containing the full student data.
        subject (str): The subject that was updated.
        new_marks (str): The new marks value.
        faculty_name (str): Name of the faculty who made the change.
        save_dir (str): Directory to save the Word file (default is current directory).
    """
    try:
        # Create file name:
        file_name = f"{student['Name'].replace(' ', '')}{student['Enrollment']}.docx"
        file_path = os.path.join(save_dir, file_name)

        # Timestamp
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M")

        # Standard order of fields
        ordered_fields = ["Roll", "Name", "Enrollment", "Password", "Sex", "DSA", "AC", "DC", "MATLAB"]

        if os.path.exists(file_path):
            doc = Document(file_path)

            if doc.paragraphs:
                # Append to existing paragraph
                last_para = doc.paragraphs[-1]
                update_entry = f"{subject}(updated at {timestamp} by {faculty_name}): {new_marks}"
                if update_entry not in last_para.text:
                    last_para.text += f", {update_entry}"
            else:
                # Edge case: file exists but is empty — write fresh full line + update
                full_line = ", ".join([f"{key}: {student[key]}" for key in ordered_fields if key in student])
                para = doc.add_paragraph(full_line)
                para.text += f", {subject}(updated at {timestamp} by {faculty_name}): {new_marks}"

        else:
            # Create new file and write full record
            doc = Document()
            full_line = ", ".join([f"{key}: {student[key]}" for key in ordered_fields if key in student])
            para = doc.add_paragraph(full_line)
            para.text += f", {subject}(updated at {timestamp} by {faculty_name}): {new_marks}"

        doc.save(file_path)
        print(f"Word file updated: {file_path}")

    except Exception as e:
        print(f"Error updating Word file for {student['Enrollment']}: {e}")
